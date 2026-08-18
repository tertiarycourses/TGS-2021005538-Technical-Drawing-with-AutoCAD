#!/usr/bin/env python3
"""Generate the WSQ Technical Drawing with AutoCAD Lesson Plan (LP) DOCX in the
Tertiary house format.

Cover page + Document Version Control Record + auto TOC + Arial 11pt body +
colour-coded 2-day schedule tables (9:30am-6:30pm, 8 training hours/day, 1h
lunch, tea within, final assessment Day 2 from 4:00pm: WA 60 min + PP 90 min).
Topics/labs come from course_data + the domain data files, and every schedule
row carries the slide numbers of the current deck (slide_marks.json, written by
build_slides.py) so the LP stays aligned with the deck, guide and labs.
"""
import os, sys, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

MARKS=json.load(open(os.path.join(HERE,"slide_marks.json")))
def SL(a,b):
    """Slide range text from two mark keys (end exclusive of mark b)."""
    return f"Slides {MARKS[a]}–{MARKS[b]-1}"

BRAND=RGBColor(0x1F,0x6F,0xEB); GREY=RGBColor(0x55,0x5B,0x66)
HEADER_FILL="1F6FEB"; TOPIC_FILL="E8F0FE"; BREAK_FILL="FFF4E5"; LUNCH_FILL="FDE9D9"; ASSESS_FILL="E8F7EE"

def lab_titles(nums):
    return "; ".join(f"Lab {a['num']}: {a['title']}" for a in ACT if a['num'] in nums)

# ------------------------------------------------ schedule (single source of truth for timing)
# (start, end, minutes, kind, activity_text, slide_ref)
SCHEDULE = {
 1: (C.DAY_THEMES[1], [
    ("9:30","10:00",30,"admin","Welcome, course introduction, ground rules, trainer and learner introductions, learning outcomes, course outline and mandatory digital attendance (AM)", SL("attendance_front","topic1")),
    ("10:00","10:45",45,"topic","Topic 1 — Getting Started with AutoCAD: user interface, workspaces, model/paper space, navigation, precision aids, units (concepts + demo)", SL("topic1","lab1")),
    ("10:45","11:15",30,"lab","Hands-on: "+lab_titles([1]), SL("lab1","topic2")),
    ("11:15","11:30",15,"break","Tea break",""),
    ("11:30","12:30",60,"topic","Topic 2 — Technical Drawing with AutoCAD: lines, circles, rectangles, polygons, ellipses, arcs, polylines, splines, construction geometry, isometric drawing (concepts + demo)", SL("topic2","lab2")),
    ("12:30","13:30",60,"lunch","Lunch break",""),
    ("13:30","15:30",120,"lab","Hands-on: "+lab_titles([2,3,4,5]), SL("lab2","topic3")),
    ("15:30","15:45",15,"break","Tea break",""),
    ("15:45","16:30",45,"topic","Topic 3 — Modifying Objects: move/copy/rotate/offset, mirror/stretch/scale, trim/extend, fillet/chamfer, arrays (concepts + demo)", SL("topic3","lab6")),
    ("16:30","18:15",105,"lab","Hands-on: "+lab_titles([6,7,8]), SL("lab6","topic4")),
    ("18:15","18:30",15,"recap","Day 1 recap, Q&A and PM digital attendance",""),
 ]),
 2: (C.DAY_THEMES[2], [
    ("9:30","9:45",15,"recap","Day 1 recap and mandatory digital attendance (AM)",""),
    ("9:45","10:25",40,"topic","Topic 4 — Organizing Objects: object properties, layers, grips, hide/isolate, blocks & attributes, external references (concepts + demo)", SL("topic4","lab9")),
    ("10:25","11:10",45,"lab","Hands-on: "+lab_titles([9,10]), SL("lab9","lab11")),
    ("11:10","11:25",15,"break","Tea break",""),
    ("11:25","12:15",50,"lab","Hands-on: "+lab_titles([11,12]), SL("lab11","topic5")),
    ("12:15","13:15",60,"lunch","Lunch break",""),
    ("13:15","13:55",40,"topic","Topic 5 — Design Documentation: text & styles, dimensions, multileaders, tables & part lists, layouts, viewports, plotting, hatching (concepts + demo)", SL("topic5","lab13")),
    ("13:55","14:40",45,"lab","Hands-on: "+lab_titles([13,14]), SL("lab13","lab15")),
    ("14:40","14:55",15,"break","Tea break",""),
    ("14:55","15:30",35,"lab","Hands-on: "+lab_titles([15,16]), SL("lab15","wrapup")),
    ("15:30","15:45",15,"recap","Course recap, feedback, TRAQOM survey and practice exam briefing", SL("wrapup","assessment_end")),
    ("15:45","16:00",15,"assess","Briefing for Assessment and Assessment digital attendance",""),
    ("16:00","17:00",60,"assess","Written Assessment (WA) — Short-Answer Questions (SAQ), 60 minutes, open book",""),
    ("17:00","18:30",90,"assess","Practical Performance (PP) — produce a dimensioned AutoCAD technical drawing, 90 minutes, open book. PM digital attendance",""),
 ]),
}

# ------------------------------------------------ build document
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)

prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("11",C.VERSION_DATE,"Legacy release — master trainer slides v11 with 2-day schedule.",C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "Full regeneration in the single-source pipeline: deck rebuilt in the house visual style (151 slides), "
  "16 hands-on labs in per-lab folders (DWG + PDF + Markdown), schedule re-timed to 9:30am–6:30pm with "
  "WA 60 min + PP 90 min on Day 2, and slide references added to every schedule row.",C.TRAINER),
])
prodoc.add_toc(doc)

def H(text,level=1):
    return doc.add_heading(text,level=level)

H("Course Information",1)
info=[("Course Title",C.TITLE),("WSQ Course Reference",C.COURSE_CODE),
      ("Skills Framework TSC",f"{C.TSC_TITLE} ({C.TSC_CODE})"),
      ("Training Provider",C.ORG+"  ("+C.UEN.replace('UEN: ','UEN ')+")"),
      ("Duration","2 days · 8 training hours per day (16 hours, including assessment)"),
      ("Daily Timing","9:30 am – 6:30 pm (1-hour lunch; tea breaks within training time)"),
      ("Mode","Instructor-led, hands-on AutoCAD drafting labs — one computer with AutoCAD per learner"),
      ("Trainer",C.TRAINER)]
t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
for k,v in info:
    c=t.add_row().cells; c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
    prodoc._shade_cell(c[0],TOPIC_FILL)
    c[1].text=""; c[1].paragraphs[0].add_run(v).font.size=Pt(10)

H("Learning Outcomes",1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for lo in C.LEARNING_OUTCOMES:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(lo).font.size=Pt(10.5)

H("Assessment",1)
for a in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],
          "Format: Open Book — course slides, Learner Guide and approved materials only.",
          "Final assessment is conducted on Day 2 from 4:00 pm (WA 60 min, PP 90 min).",
          C.ASSESSMENT["note"]]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10.5)

def set_cell(cell,text,bold=False,size=9.5,color=None,fill=None):
    cell.text=""; p=cell.paragraphs[0]
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name="Arial"
    if color: r.font.color.rgb=color
    if fill: prodoc._shade_cell(cell,fill)

KIND_FILL={"topic":TOPIC_FILL,"break":BREAK_FILL,"lunch":LUNCH_FILL,"assess":ASSESS_FILL,
           "admin":"F3F5F8","recap":"F3F5F8","lab":None}

H("Course Schedule",1)
for day,(theme,rows) in SCHEDULE.items():
    H(f"Day {day} — {theme}",2)
    tbl=doc.add_table(rows=0,cols=4); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=tbl.add_row().cells
    for i,htext in enumerate(["Time","Duration","Topic / Activity","Slides"]):
        set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
    training=0
    for start,end,mins,kind,text,slides in rows:
        cells=tbl.add_row().cells; fill=KIND_FILL.get(kind)
        set_cell(cells[0],f"{start}–{end}",bold=(kind in ("topic","assess")),fill=fill)
        set_cell(cells[1],f"{mins} min",fill=fill)
        set_cell(cells[2],text,bold=(kind in ("topic","assess")),fill=fill)
        set_cell(cells[3],slides,size=9,fill=fill)
        if kind!="lunch": training+=mins
    for row in tbl.rows:
        row.cells[0].width=Inches(1.05); row.cells[1].width=Inches(0.85)
        row.cells[2].width=Inches(4.05); row.cells[3].width=Inches(1.05)
    p=doc.add_paragraph(); r=p.add_run(f"Total training time: {training} minutes ({training//60} hours)."); r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
    assert training==480, f"Day {day} training minutes = {training}, expected 480"

H("Lab Reference (aligned to the Skills Framework TSC)",1)
tt=doc.add_table(rows=0,cols=3); tt.style="Table Grid"
hdr=tt.add_row().cells
for i,htext in enumerate(["Topic / TSC mapping","Ability · Knowledge","Labs"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    cells=tt.add_row().cells
    set_cell(cells[0],f"Topic {tp['code']}: {tp['title']}",bold=True,fill=TOPIC_FILL)
    set_cell(cells[1],tp["weighting"],fill=TOPIC_FILL)
    set_cell(cells[2],", ".join(f"Lab {a['num']}" for a in acts))

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
doc.save(OUT)
print("Saved",OUT)

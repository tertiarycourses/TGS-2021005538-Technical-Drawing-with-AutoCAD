#!/usr/bin/env python3
"""Generate the WSQ Technical Drawing with AutoCAD Learner Guide as BOTH a
Markdown mirror (LG-*.md at repo root) and a DOCX (courseware/LG-*.docx) from
one source, so they never diverge.

House format: cover page, Document Version Control Record, auto TOC, Arial 11pt
body, one section per lab (Objective · Goal · You'll produce · Step-by-step ·
Check your work), plus setup, drawing-reading guidance and glossary. All content
is driven by course_data + the domain data files, keeping the LG 100% aligned
with the slide deck, Lesson Plan and labs.
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
from labslug import lab_dirname
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

# ---------------- block DSL (single content stream → MD + DOCX) ----------------
B=[]
def h1(t): B.append(("h1",t))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def p(t):  B.append(("p",t))
def bullets(xs): B.append(("bullets",xs))
def steps(xs): B.append(("steps",xs))
def code(t): B.append(("code",t))
def note(t): B.append(("note",t))
def img(path,caption): B.append(("img",path,caption))
def rule(): B.append(("rule",))

# ---------------- content ----------------
h1("Introduction")
p(f"This Learner Guide accompanies the WSQ course {C.TITLE} ({C.COURSE_CODE}), conducted by {C.ORG}. "
  f"It provides step-by-step instructions for all 16 hands-on labs, organised by the five course topics, "
  f"aligned to the Skills Framework TSC {C.TSC_TITLE} ({C.TSC_CODE}).")
p("Use this guide alongside the course slides and the lab folders in labs/ of the course repository. "
  "Each lab folder contains this guide's instructions as README.md, a printable PDF instruction sheet, "
  "and the AutoCAD DWG data files for that lab. The final assessment is open book: you may refer to "
  "the slides, this Learner Guide and other approved materials.")

h1("Course Learning Outcomes")
bullets(C.LEARNING_OUTCOMES)

h1("Before You Start — Environment Setup")
h3("What you need")
bullets([
 "A Windows PC (or Mac) with AutoCAD installed — AutoCAD 2021 or later (any recent release works; the free 30-day trial or an education licence is fine).",
 "Download AutoCAD from https://www.autodesk.com/products/autocad/free-trial or the education community at https://www.autodesk.com/education.",
 "The course lab files — the labs/ folder of the course repository (one folder per lab: DWG data files, PDF instruction sheet and README).",
 "A mouse with a scroll wheel — zooming and panning with the wheel is essential for CAD work.",
])
h3("Set up AutoCAD for the course")
bullets([
 "Start AutoCAD and choose the Drafting & Annotation workspace.",
 "New drawings in this course start from the metric template acadiso.dwt.",
 "Keep the status bar visible — grid, snap, ortho, polar tracking and object snap are used in every lab.",
 "Copy the labs/ folder to your Documents so you can open and save the lab drawings freely.",
])
h3("Conventions used in every lab")
bullets([
 "Commands are typed at the command line or picked from the ribbon; the short alias is shown in brackets, e.g. LINE (L).",
 "The → arrow chains a command's prompts, e.g. 'C → D → 40' means run CIRCLE, choose Diameter, type 40.",
 "Coordinates: absolute 100,50 · relative @50,0 · polar @50<30.",
 "Save your work as <lab>-mine.dwg so the supplied solution files stay untouched.",
 "Every lab ends with a 'Check your work' test — verify before moving on.",
])

# ---------------- per-topic, per-lab ----------------
for t in C.TOPICS:
    h1(f"Topic {t['code']} — {t['title']}  ({t['weighting']})")
    p(t["subtitle"])
    h3("Key concepts")
    bullets([f"{head} — {body}" for head,body in t["concepts"]])
    for a in [x for x in ACT if x["topic"]==t["num"]]:
        h2(f"Lab {a['num']} — {a['title']}")
        p(f"Objective: {a['objective']}.")
        p(f"Goal: {a['desc']}")
        if a.get("img"):
            img(f"labs/{lab_dirname(a)}/{a['img']}", f"Lab {a['num']} — the drawing you'll produce")
        h3("You'll produce")
        p(a["build"]+f"   (Commands: {a['services']}.)")
        h3("Step-by-step")
        steps(list(a["steps"]))
        h3("Check your work")
        p(a["test"])
        note(f"This lab's data files and printable instructions are in labs/{lab_dirname(a)}/.")
        rule()

h1("Reading Technical Drawings — Quick Reference")
p("These conventions (K1/K3) apply across every lab and are assessed in the written assessment.")
bullets([
 "Line conventions: continuous thick = visible outline; dashed = hidden detail; chain thin = centreline; continuous thin = dimensions, hatching and projections.",
 "Orthographic projection presents related 2D views (front, top, side) aligned by projection; third-angle and first-angle layouts differ by where the top/side views sit.",
 "Isometric drawings show three faces at once with receding axes at 30° — a pictorial view, not a true 3D model.",
 "Section views cut through the object; the cut material is hatched (e.g. ANSI31) with adjacent parts hatched at different angles.",
 "Dimensions state real object size regardless of plotted scale; never measure a printed sheet with a ruler.",
 "The title block records the drawing title, number, scale, units, author, date, revision and sheet size.",
])

h1("Preparing for the Assessment")
bullets([
 "First pass: complete every lab following this guide, checking each 'Check your work' test.",
 "Second pass: redo Labs 2–8 from just the drawing-sheet PDFs, without the step-by-step instructions.",
 "Review the Key Concepts bullets of each topic — the written assessment (SAQ) draws on them.",
 "The practical performance asks you to produce a dimensioned technical drawing end-to-end: set up, draw, modify, organise with layers, dimension, and plot — Labs 2, 7, 9, 14 and 16 cover that chain.",
 f"Course materials, assessment submission and your certificate are on the LMS: {C.LMS_URL}",
])

h1("Glossary")
gl=[
 ("Model space","The unlimited drawing environment where designs are drawn at full (1:1) size."),
 ("Paper space / Layout","The sheet-composition environment where scaled viewports, title blocks and notes are arranged for printing."),
 ("Viewport","A window on a layout that displays a view of model space at a set scale."),
 ("Workspace","A task-oriented arrangement of ribbons, menus and palettes (e.g. Drafting & Annotation)."),
 ("Object snap (Osnap)","Precision aid that picks exact geometric points — endpoints, midpoints, centres, intersections."),
 ("Polar tracking","Guides the cursor along preset angle increments while drawing."),
 ("Polyline","A single object made of connected line and arc segments."),
 ("Isometric drawing","A flat 2D drawing that represents a 3D object using axes at 30°/90°/150° (ISODRAFT)."),
 ("Layer","A named organising level that controls the colour, linetype, lineweight, visibility and plottability of the objects on it."),
 ("Block","A named, reusable group of objects inserted as one unit; dynamic blocks add adjustable parameters."),
 ("Attribute","A text data tag attached to a block — part number, description, quantity — extractable into tables."),
 ("External reference (Xref)","Another drawing attached live into the current drawing rather than copied in."),
 ("Dimension style","A named collection of dimension settings (arrows, text, units) applied consistently across a drawing."),
 ("Multileader","An annotation object combining an arrowhead, leader line, landing and note text."),
 ("Hatch","A pattern, solid or gradient fill applied to an enclosed area, e.g. to show cut material in sections."),
 ("Plot","Output a drawing to a printer/plotter or to a file such as PDF; plotting is configured per layout."),
]
B.append(("dl",gl))

# ---------------- render Markdown ----------------
def _anchor(txt):
    return "".join(ch.lower() if ch.isalnum() else ("-" if ch in " -" else "") for ch in txt)

def render_md():
    out=[f"# {C.TITLE} — Learner Guide",""]
    out.append(f"**WSQ Course Code:** {C.COURSE_CODE}  |  **Conducted by:** {C.ORG} ({C.UEN.replace('UEN: ','UEN ')})  |  **Version {C.VERSION} · {C.VERSION_DATE}**")
    out.append("")
    out.append("## Contents"); out.append("")
    for kind,*rest in B:
        if kind=="h1": out.append(f"- [{rest[0]}](#{_anchor(rest[0])})")
        elif kind=="h2": out.append(f"  - [{rest[0]}](#{_anchor(rest[0])})")
    out.append("")
    for kind,*rest in B:
        if kind=="h1": out+=["",f"## {rest[0]}",""]
        elif kind=="h2": out+=["",f"### {rest[0]}",""]
        elif kind=="h3": out+=[f"**{rest[0]}**",""]
        elif kind=="p": out+=[rest[0],""]
        elif kind=="bullets": out+=[f"- {x}" for x in rest[0]]+[""]
        elif kind=="steps":
            for i,(instr,cmd) in enumerate(rest[0],1):
                out.append(f"{i}. {instr}")
                if cmd: out+=["","   ```",f"   {cmd}","   ```",""]
            out.append("")
        elif kind=="code": out+=["```",rest[0],"```",""]
        elif kind=="img": out+=[f"![{rest[1]}]({rest[0]})","",f"*{rest[1]}*",""]
        elif kind=="note": out+=[f"> **Note:** {rest[0]}",""]
        elif kind=="rule": out+=["---",""]
        elif kind=="dl":
            for term,defn in rest[0]: out.append(f"- **{term}** — {defn}")
            out.append("")
    return "\n".join(out)

MD_OUT=os.path.join(REPO,f"LG-{C.SHORT_TITLE}.md")
with open(MD_OUT,"w") as f: f.write(render_md())
print("Saved",MD_OUT)

# ---------------- render DOCX ----------------
BRAND=RGBColor(0x1F,0x6F,0xEB); GREY=RGBColor(0x55,0x5B,0x66)
INKCODE=RGBColor(0x0B,0x30,0x60)
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)
prodoc.add_cover_page(doc,"LEARNER GUIDE",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("11","18 August 2026","Legacy release — learner guide aligned to master trainer slides v11.",C.TRAINER),
 ("12","18 August 2026",
  "Full single-source regeneration with 16 step-by-step labs, drawing-reading reference, "
  "assessment preparation guidance and glossary.",C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "DWG integration release: strengthened Labs 1, 8–12, 15 and 16 with active practice files, "
  "repaired the Lab 12 Landscaping_REV1 relative Xref dependency, and retained all detailed "
  "procedures in the Learner Guide and lab packs for alignment with the concept-led v13 deck.",C.TRAINER),
])
prodoc.add_toc(doc)

def code_para(text):
    for line in text.split("\n"):
        para=doc.add_paragraph()
        para.paragraph_format.keep_together=True
        r=para.add_run(line); r.font.name="Consolas"; r.font.size=Pt(9.5); r.font.color.rgb=INKCODE

for kind,*rest in B:
    if kind=="h1":
        para=doc.add_heading(rest[0],level=1); para.paragraph_format.keep_with_next=True
    elif kind=="h2":
        para=doc.add_heading(rest[0],level=2); para.paragraph_format.keep_with_next=True
    elif kind=="h3":
        para=doc.add_paragraph(); r=para.add_run(rest[0]); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=BRAND
        para.paragraph_format.keep_with_next=True
    elif kind=="p": doc.add_paragraph(rest[0])
    elif kind=="bullets":
        for x in rest[0]: doc.add_paragraph(x,style="List Bullet")
    elif kind=="steps":
        for i,(instr,cmd) in enumerate(rest[0],1):
            para=doc.add_paragraph(); para.paragraph_format.left_indent=Inches(0.25)
            para.paragraph_format.keep_together=True
            para.paragraph_format.keep_with_next=bool(cmd)
            r=para.add_run(f"{i}.  "); r.bold=True
            para.add_run(instr)
            if cmd: code_para("     "+cmd)
    elif kind=="code": code_para(rest[0])
    elif kind=="img":
        ipath=os.path.join(REPO,rest[0])
        if os.path.exists(ipath):
            try:
                doc.add_picture(ipath,width=Inches(4.2))
                cap=doc.add_paragraph(); rc_=cap.add_run(rest[1]); rc_.italic=True; rc_.font.size=Pt(9); rc_.font.color.rgb=GREY
            except Exception as e:
                print("WARN image skip:",rest[0],e)
    elif kind=="note":
        para=doc.add_paragraph(); r=para.add_run("Note: "); r.bold=True; r.font.color.rgb=BRAND
        para.add_run(rest[0]).font.size=Pt(10)
    elif kind=="rule": doc.add_paragraph("")
    elif kind=="dl":
        for term,defn in rest[0]:
            para=doc.add_paragraph(style="List Bullet")
            r=para.add_run(term+" — "); r.bold=True; para.add_run(defn)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
DOCX_OUT=os.path.join(REPO,"courseware",f"LG-{C.SHORT_TITLE}.docx")
doc.save(DOCX_OUT)
print("Saved",DOCX_OUT)
